"""It turns ANY supported file into plain text."""
import csv
import logging
import os
import tempfile
import uuid
import zipfile
from email import message_from_bytes
from email.policy import default as default_email_policy
from typing import Any, Dict, Iterator, Optional

import fitz  # PyMuPDF. The package installs as `pymupdf` but imports as `fitz`.
import openpyxl
from bs4 import BeautifulSoup
from docx import Document as DocxDocument
from langdetect import detect
from pptx import Presentation

from app.core.config import settings
from app.services.ocr import OCRService

logger = logging.getLogger(__name__)

IMAGE_EXTENSIONS = frozenset({"png", "jpg", "jpeg", "tiff", "bmp", "gif"})
SUPPORTED_EXTENSIONS = frozenset(
    {
        "pdf", "docx", "pptx", "xlsx", "xls", "html", "htm",
        "eml", "txt", "md", "markdown", "csv", "zip",
    }
    | IMAGE_EXTENSIONS
)

OFFICE_ZIP_EXTENSIONS = frozenset({"docx", "pptx", "xlsx"})


class UnsafeArchiveError(ValueError):
    """A zip archive broke one of the ZIP_* safety limits."""


class _ZipBudget:
    """The remaining allowance for one upload, shared by nested archives."""

    def __init__(self):
        self.entries_left = settings.ZIP_MAX_ENTRIES
        self.bytes_left = settings.ZIP_MAX_TOTAL_UNCOMPRESSED_BYTES

    def take_entry(self) -> None:
        self.entries_left -= 1
        if self.entries_left < 0:
            raise UnsafeArchiveError(
                f"Archive rejected: more than {settings.ZIP_MAX_ENTRIES} files."
            )

    def take_bytes(self, count: int) -> None:
        self.bytes_left -= count
        if self.bytes_left < 0:
            raise UnsafeArchiveError(
                "Archive rejected: it expands to more than "
                f"{settings.ZIP_MAX_TOTAL_UNCOMPRESSED_BYTES // (1024 * 1024)} MB."
            )


class DocumentTooLarge(ValueError):
    """The document is over a per-document cap (INGESTION_MAX_PAGES or"""


class Segment:
    """One piece of a document, produced while the file is still being read."""

    __slots__ = ("page_number", "text", "ocr", "meta")

    def __init__(self, page_number: Optional[int], text: str, ocr: bool = False, meta: Optional[dict] = None):
        self.page_number = page_number
        self.text = text
        self.ocr = ocr
        self.meta = meta or {}


class ExtractionResult:
    """The uniform result every extractor returns, whatever the input format."""

    def __init__(
        self,
        text: str,
        metadata: Dict[str, Any],
        language: str,
        needs_ocr: bool = False,
    ):
        self.text = text
        self.metadata = metadata
        self.language = language
        self.needs_ocr = needs_ocr


class DocumentExtractor:
    """Routes each file type to its dedicated parser."""


    @staticmethod
    def detect_language(text: str) -> str:
        """Work out which human language the text is written in."""
        if not text or len(text.strip()) < 10:
            return "en"
        try:
            return detect(text)
        except Exception:
            return "en"


    @classmethod
    def extract(cls, file_path: str, file_type: str) -> ExtractionResult:
        """Send the file to whichever parser handles its extension."""
        file_type = file_type.lower()

        if file_type in OFFICE_ZIP_EXTENSIONS:
            cls._check_office_container(file_path)

        if file_type == "pdf":
            return cls.extract_pdf(file_path)
        elif file_type in IMAGE_EXTENSIONS:
            return cls.extract_image(file_path, file_type)
        elif file_type == "docx":
            return cls.extract_docx(file_path)
        elif file_type == "pptx":
            return cls.extract_pptx(file_path)
        elif file_type in ["xlsx", "xls"]:
            return cls.extract_xlsx(file_path)
        elif file_type in ["html", "htm"]:
            return cls.extract_html(file_path)
        elif file_type == "eml":
            return cls.extract_eml(file_path)
        elif file_type in ["txt", "md", "markdown"]:
            return cls.extract_txt_or_md(file_path)
        elif file_type == "csv":
            return cls.extract_csv(file_path)
        elif file_type == "zip":
            return cls.extract_zip(file_path)
        else:
            raise ValueError(f"Unsupported file type: {file_type}")


    _TEXT_SEGMENT_BYTES = 256 * 1024

    @classmethod
    def iter_segments(cls, file_path: str, file_type: str, info: Dict[str, Any]) -> Iterator[Segment]:
        """Yield the document a Segment at a time."""
        file_type = file_type.lower()
        if file_type in OFFICE_ZIP_EXTENSIONS:
            cls._check_office_container(file_path)

        if file_type == "pdf":
            yield from cls._iter_pdf(file_path, info)
        elif file_type in IMAGE_EXTENSIONS:
            yield from cls._iter_image(file_path, file_type, info)
        elif file_type == "csv":
            yield from cls._iter_csv(file_path, info)
        elif file_type == "xlsx":
            yield from cls._iter_xlsx(file_path, info)
        elif file_type in ("txt", "md", "markdown"):
            yield from cls._iter_text(file_path, info)
        else:
            result = cls.extract(file_path, file_type)
            info.update(result.metadata)
            if result.needs_ocr:
                info["needs_ocr"] = True
            yield Segment(None, result.text, ocr=result.needs_ocr)


    @staticmethod
    def _image_coverage(page) -> float:
        """What fraction of the page's area is covered by pictures (0..1)."""
        try:
            page_area = abs(page.rect) or 1.0
            covered = 0.0
            for image in page.get_image_info():
                box = fitz.Rect(image["bbox"]) & page.rect
                covered += abs(box)
            return min(1.0, covered / page_area)
        except Exception:
            return 0.0

    @classmethod
    def _iter_pdf(cls, file_path: str, info: Dict[str, Any]) -> Iterator[Segment]:
        """Yield each page's text, OCR'ing the pages that need it."""
        from app.services import vision

        document = fitz.open(file_path)
        try:
            page_count = len(document)
            info.update(
                {
                    "title": document.metadata.get("title", "") if document.metadata else "",
                    "author": document.metadata.get("author", "") if document.metadata else "",
                    "creator": document.metadata.get("creator", "") if document.metadata else "",
                    "page_count": page_count,
                }
            )
            if page_count > settings.INGESTION_MAX_PAGES:
                raise DocumentTooLarge(
                    f"This PDF has {page_count:,} pages; the limit is "
                    f"{settings.INGESTION_MAX_PAGES:,} (INGESTION_MAX_PAGES). "
                    "Split it into smaller files, or raise the limit."
                )
            captions_left = settings.VISION_MAX_IMAGES_PER_DOCUMENT if vision.enabled() else 0
            ocr_pages = 0
            captions = 0

            for index in range(page_count):
                page = document.load_page(index)
                text = page.get_text()
                used_ocr = False
                if len(text.strip()) < settings.OCR_CHAR_THRESHOLD_PER_PAGE and (
                    page.get_images() or page.get_cdrawings()
                ):
                    ocr_text = OCRService.ocr_page(page)
                    text = f"{text.strip()}\n{ocr_text}".strip()
                    used_ocr = True
                    ocr_pages += 1
                if captions_left > 0 and cls._image_coverage(page) >= settings.VISION_MIN_IMAGE_COVERAGE:
                    caption = vision.caption_image(OCRService.render_page(page))
                    if caption:
                        text += vision.caption_block(caption)
                        captions += 1
                    captions_left -= 1
                del page
                if index % 25 == 24:
                    fitz.TOOLS.store_shrink(100)
                yield Segment(index + 1, text, ocr=used_ocr)

            if ocr_pages:
                info.update(
                    {"ocr_applied": True, "ocr_provider": settings.OCR_PROVIDER, "ocr_pages": ocr_pages, "needs_ocr": True}
                )
            if captions:
                info["vision_captions"] = captions
        finally:
            document.close()
            fitz.TOOLS.store_shrink(100)


    @classmethod
    def _iter_image(cls, file_path: str, file_type: str, info: Dict[str, Any]) -> Iterator[Segment]:
        """OCR the image, add a caption when vision is on, yield it once."""
        from app.services import vision

        result = cls.extract_image(file_path, file_type)
        info.update(result.metadata)
        info["needs_ocr"] = True
        text = result.text
        if vision.enabled() and "error" not in result.metadata:
            with open(file_path, "rb") as handle:
                caption = vision.caption_image(handle.read())
            if caption:
                text += vision.caption_block(caption)
                info["vision_captions"] = 1
        yield Segment(1, text, ocr=True)


    @staticmethod
    def _row_segment(title: str, header: Optional[str], rows: list, first_row: int, meta: dict) -> Segment:
        """A batch of rows as one segment. The column header is repeated at the"""
        last_row = first_row + len(rows) - 1
        lines = [f"{title} (rows {first_row}-{last_row}):"] if title else []
        if header is not None and first_row > 1:
            lines.append(header)
        lines.extend(rows)
        return Segment(1, "\n".join(lines), meta={**meta, "rows": f"{first_row}-{last_row}"})

    @classmethod
    def _iter_csv(cls, file_path: str, info: Dict[str, Any]) -> Iterator[Segment]:
        """A CSV, INGESTION_ROWS_PER_SEGMENT rows at a time."""
        per_segment = max(1, settings.INGESTION_ROWS_PER_SEGMENT)
        total = 0
        try:
            with open(file_path, "r", encoding="utf-8", errors="ignore", newline="") as handle:
                reader = csv.reader(handle)
                header: Optional[str] = None
                batch: list = []
                first = 1
                for row in reader:
                    line = " , ".join(row)
                    total += 1
                    if header is None:
                        header = line
                    batch.append(line)
                    if len(batch) >= per_segment:
                        yield cls._row_segment("", header, batch, first, {})
                        first += len(batch)
                        batch = []
                if batch:
                    yield cls._row_segment("", header, batch, first, {})
        except Exception as exc:
            yield Segment(1, f"CSV extraction failed: {exc}")
        info["row_count"] = total

    @classmethod
    def _iter_xlsx(cls, file_path: str, info: Dict[str, Any]) -> Iterator[Segment]:
        """A workbook, sheet by sheet, INGESTION_ROWS_PER_SEGMENT rows at a time."""
        per_segment = max(1, settings.INGESTION_ROWS_PER_SEGMENT)
        try:
            workbook = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
        except Exception as exc:
            yield Segment(1, f"Excel extraction failed: {exc}")
            return
        try:
            info["sheet_names"] = list(workbook.sheetnames)
            for sheet_name in workbook.sheetnames:
                sheet = workbook[sheet_name]
                header: Optional[str] = None
                batch: list = []
                first = 1
                for row in sheet.iter_rows(values_only=True):
                    if not any(value is not None and value != "" for value in row):
                        continue
                    line = " | ".join("" if value is None else str(value) for value in row)
                    if header is None:
                        header = line
                    batch.append(line)
                    if len(batch) >= per_segment:
                        yield cls._row_segment(f"Sheet '{sheet_name}'", header, batch, first, {"sheet": sheet_name})
                        first += len(batch)
                        batch = []
                if batch:
                    yield cls._row_segment(f"Sheet '{sheet_name}'", header, batch, first, {"sheet": sheet_name})
        finally:
            workbook.close()


    @classmethod
    def _iter_text(cls, file_path: str, info: Dict[str, Any]) -> Iterator[Segment]:
        """A text or Markdown file in ~256 KB slices, cut at line breaks."""
        try:
            with open(file_path, "r", encoding="utf-8", errors="ignore") as handle:
                buffer: list = []
                size = 0
                for line in handle:
                    buffer.append(line)
                    size += len(line)
                    if size >= cls._TEXT_SEGMENT_BYTES:
                        yield Segment(1, "".join(buffer))
                        buffer, size = [], 0
                if buffer:
                    yield Segment(1, "".join(buffer))
        except Exception as exc:
            yield Segment(1, f"Text extraction failed: {exc}")


    @classmethod
    def extract_pdf(cls, file_path: str) -> ExtractionResult:
        """Extract text from a PDF, falling back to OCR when it is a scan."""
        text = ""
        metadata: Dict[str, Any] = {}
        needs_ocr = False

        try:
            document = fitz.open(file_path)

            metadata = {
                "title": document.metadata.get("title", ""),
                "author": document.metadata.get("author", ""),
                "creator": document.metadata.get("creator", ""),
                "page_count": len(document),
            }

            pages_text = [page.get_text() for page in document]
            text = "\n--- PAGE BREAK ---\n".join(pages_text)

            page_count = len(document)
            threshold = settings.OCR_CHAR_THRESHOLD_PER_PAGE * page_count

            if page_count > 0 and len(text.strip()) < threshold:
                logger.info(
                    f"Only {len(text.strip())} characters across "
                    f"{page_count} pages -- below the {threshold} threshold. "
                    f"Treating as a scan and running OCR."
                )
                ocr_text, _ = OCRService.ocr_pdf(file_path)
                text = ocr_text
                needs_ocr = True
                metadata["ocr_applied"] = True
                metadata["ocr_provider"] = settings.OCR_PROVIDER

            document.close()

        except Exception as exc:
            logger.warning(f"PDF parsing failed: {exc}. Attempting OCR...")
            try:
                ocr_text, _ = OCRService.ocr_pdf(file_path)
                text = ocr_text
                needs_ocr = True
                metadata["ocr_applied"] = True
                metadata["ocr_provider"] = settings.OCR_PROVIDER
                metadata["parse_error"] = str(exc)
            except Exception as ocr_error:
                text = f"PDF extraction and OCR fallback both failed: {ocr_error}"
                needs_ocr = True
                metadata["error"] = str(ocr_error)

        return ExtractionResult(
            text=text,
            metadata=metadata,
            language=cls.detect_language(text),
            needs_ocr=needs_ocr,
        )


    @classmethod
    def extract_image(cls, file_path: str, file_type: str) -> ExtractionResult:
        """Read the text out of a photograph or screenshot."""
        try:
            with open(file_path, "rb") as handle:
                image_bytes = handle.read()

            text = OCRService.ocr_image(image_bytes)
            metadata = {
                "image_ocr": True,
                "ocr_provider": settings.OCR_PROVIDER,
                "image_format": file_type,
            }
        except Exception as exc:
            text = f"Image OCR extraction failed: {exc}"
            metadata = {"error": str(exc)}

        return ExtractionResult(
            text=text,
            metadata=metadata,
            language=cls.detect_language(text),
            needs_ocr=True,
        )


    @classmethod
    def extract_docx(cls, file_path: str) -> ExtractionResult:
        """Extract a Word document's paragraphs AND its tables."""
        text = ""
        metadata: Dict[str, Any] = {}

        try:
            document = DocxDocument(file_path)

            paragraphs = [p.text for p in document.paragraphs]

            table_rows = []
            for table in document.tables:
                for row in table.rows:
                    cells = [cell.text.strip() for cell in row.cells]
                    table_rows.append(" | ".join(cells))

            text = "\n".join(paragraphs)
            if table_rows:
                text += "\n\nTables:\n" + "\n".join(table_rows)

            metadata = {
                "paragraph_count": len(paragraphs),
                "table_count": len(document.tables),
            }
        except Exception as exc:
            text = f"DOCX extraction failed: {exc}"
            metadata = {"error": str(exc)}

        return ExtractionResult(
            text=text, metadata=metadata, language=cls.detect_language(text)
        )


    @classmethod
    def extract_pptx(cls, file_path: str) -> ExtractionResult:
        """Pull text from every shape on every slide."""
        text = ""
        metadata: Dict[str, Any] = {}

        try:
            presentation = Presentation(file_path)

            slides_text = []
            for index, slide in enumerate(presentation.slides, 1):
                shape_texts = []
                for shape in slide.shapes:
                    if hasattr(shape, "text") and shape.text.strip():
                        shape_texts.append(shape.text.strip())

                slides_text.append(f"Slide {index}:\n" + "\n".join(shape_texts))

            text = "\n\n".join(slides_text)
            metadata = {"slide_count": len(presentation.slides)}
        except Exception as exc:
            text = f"PPTX extraction failed: {exc}"
            metadata = {"error": str(exc)}

        return ExtractionResult(
            text=text, metadata=metadata, language=cls.detect_language(text)
        )


    @classmethod
    def extract_xlsx(cls, file_path: str) -> ExtractionResult:
        """Read every sheet of a workbook, row by row."""
        text = ""
        metadata: Dict[str, Any] = {}

        try:
            workbook = openpyxl.load_workbook(file_path, data_only=True)

            sheets_text = []
            for sheet_name in workbook.sheetnames:
                sheet = workbook[sheet_name]

                rows_text = []
                for row in sheet.iter_rows(values_only=True):
                    if any(row):
                        cells = [
                            str(value) if value is not None else "" for value in row
                        ]
                        rows_text.append(" | ".join(cells))

                sheets_text.append(f"Sheet '{sheet_name}':\n" + "\n".join(rows_text))

            text = "\n\n".join(sheets_text)
            metadata = {"sheet_names": workbook.sheetnames}
        except Exception as exc:
            text = f"Excel extraction failed: {exc}"
            metadata = {"error": str(exc)}

        return ExtractionResult(
            text=text, metadata=metadata, language=cls.detect_language(text)
        )


    @classmethod
    def extract_html(cls, file_path: str) -> ExtractionResult:
        """Strip the tags and keep the readable text."""
        text = ""
        metadata: Dict[str, Any] = {}

        try:
            with open(file_path, "r", encoding="utf-8", errors="ignore") as handle:
                soup = BeautifulSoup(handle.read(), "html.parser")

            raw_text = soup.get_text(separator="\n")
            text = "\n".join(
                line.strip() for line in raw_text.splitlines() if line.strip()
            )
            metadata = {"title": soup.title.string if soup.title else ""}
        except Exception as exc:
            text = f"HTML extraction failed: {exc}"
            metadata = {"error": str(exc)}

        return ExtractionResult(
            text=text, metadata=metadata, language=cls.detect_language(text)
        )


    @classmethod
    def extract_eml(cls, file_path: str) -> ExtractionResult:
        """Extract an email's headers and its plain-text body."""
        text = ""
        metadata: Dict[str, Any] = {}

        try:
            with open(file_path, "rb") as handle:
                message = message_from_bytes(
                    handle.read(), policy=default_email_policy
                )

            body = ""
            if message.is_multipart():
                for part in message.walk():
                    content_type = part.get_content_type()
                    disposition = str(part.get_params())

                    if content_type == "text/plain" and "attachment" not in disposition:
                        body += part.get_payload(decode=True).decode(errors="ignore")
            else:
                body = message.get_payload(decode=True).decode(errors="ignore")

            metadata = {
                "subject": message.get("Subject", ""),
                "from": message.get("From", ""),
                "to": message.get("To", ""),
                "date": message.get("Date", ""),
            }

            text = (
                f"From: {metadata['from']}\n"
                f"To: {metadata['to']}\n"
                f"Date: {metadata['date']}\n"
                f"Subject: {metadata['subject']}\n\n"
                f"Body:\n{body}"
            )
        except Exception as exc:
            text = f"Email extraction failed: {exc}"
            metadata = {"error": str(exc)}

        return ExtractionResult(
            text=text, metadata=metadata, language=cls.detect_language(text)
        )


    @classmethod
    def extract_txt_or_md(cls, file_path: str) -> ExtractionResult:
        """Read the file as-is. Markdown is deliberately left unrendered so the"""
        text = ""
        try:
            with open(file_path, "r", encoding="utf-8", errors="ignore") as handle:
                text = handle.read()
        except Exception as exc:
            text = f"Text extraction failed: {exc}"

        return ExtractionResult(
            text=text, metadata={}, language=cls.detect_language(text)
        )


    @classmethod
    def extract_csv(cls, file_path: str) -> ExtractionResult:
        """Flatten a CSV into readable lines, one per row."""
        text = ""
        try:
            with open(file_path, "r", encoding="utf-8", errors="ignore") as handle:
                reader = csv.reader(handle)
                rows = [" , ".join(row) for row in reader]
                text = "\n".join(rows)
        except Exception as exc:
            text = f"CSV extraction failed: {exc}"

        return ExtractionResult(
            text=text, metadata={}, language=cls.detect_language(text)
        )


    @classmethod
    def extract_zip(
        cls, file_path: str, _depth: int = 1, _budget: "_ZipBudget | None" = None
    ) -> ExtractionResult:
        """Unpack an archive and extract every supported file inside it."""
        texts = []
        metadata: Dict[str, Any] = {"files_extracted": [], "skipped_entries": []}
        budget = _budget or _ZipBudget()
        temp_dir = None

        try:
            temp_dir = tempfile.mkdtemp(prefix="docminds-zip-")

            with zipfile.ZipFile(file_path, "r") as archive:
                members = [info for info in archive.infolist() if not info.is_dir()]

                if len(members) > budget.entries_left:
                    raise UnsafeArchiveError(
                        f"Archive rejected: more than {settings.ZIP_MAX_ENTRIES} files."
                    )
                declared = sum(info.file_size for info in members)
                if declared > budget.bytes_left:
                    raise UnsafeArchiveError(
                        "Archive rejected: it expands to more than "
                        f"{settings.ZIP_MAX_TOTAL_UNCOMPRESSED_BYTES // (1024 * 1024)} MB."
                    )

                for info in members:
                    budget.take_entry()
                    display_name = os.path.basename(info.filename.replace("\\", "/"))
                    extension = os.path.splitext(display_name)[1].lstrip(".").lower()
                    reason = cls._unsafe_member_reason(info)

                    if reason is None and extension not in SUPPORTED_EXTENSIONS:
                        continue
                    if (
                        reason is None
                        and extension == "zip"
                        and _depth > settings.ZIP_MAX_NESTING_DEPTH
                    ):
                        reason = "nested archive too deep"
                    if reason is not None:
                        metadata["skipped_entries"].append(f"{info.filename} ({reason})")
                        continue

                    inner_path = os.path.join(temp_dir, f"{uuid.uuid4().hex}.{extension}")
                    try:
                        cls._copy_member(archive, info, inner_path, budget)
                        if extension == "zip":
                            result = cls.extract_zip(inner_path, _depth + 1, budget)
                        else:
                            result = cls.extract(inner_path, extension)
                        texts.append(f"--- File: {display_name} ---\n{result.text}")
                        metadata["files_extracted"].append(display_name)
                    except UnsafeArchiveError:
                        raise
                    except Exception as inner_error:
                        texts.append(
                            f"--- File: {display_name} (extraction failed: {inner_error}) ---"
                        )
                    finally:
                        try:
                            os.remove(inner_path)
                        except OSError:
                            pass

            text = "\n\n=================================\n\n".join(texts)
            metadata["file_count"] = len(metadata["files_extracted"])

        except UnsafeArchiveError:
            raise
        except Exception as exc:
            text = f"ZIP extraction failed: {exc}"
            metadata["error"] = str(exc)

        finally:
            if temp_dir and os.path.isdir(temp_dir):
                cls._remove_tree(temp_dir)

        return ExtractionResult(
            text=text, metadata=metadata, language=cls.detect_language(text)
        )

    @staticmethod
    def _check_office_container(file_path: str) -> None:
        """Refuse a .docx, .pptx or .xlsx that would expand into a memory bomb."""
        try:
            with zipfile.ZipFile(file_path) as archive:
                declared = sum(info.file_size for info in archive.infolist())
        except Exception:
            return

        limit = settings.ZIP_MAX_TOTAL_UNCOMPRESSED_BYTES
        if declared > limit:
            raise UnsafeArchiveError(
                "Document rejected: it expands to more than "
                f"{limit // (1024 * 1024)} MB when opened -- this is how zip "
                "bombs are built."
            )

    @staticmethod
    def _unsafe_member_reason(info: zipfile.ZipInfo) -> "str | None":
        """Why this member must not be extracted, or None if it is fine."""
        name = info.filename.replace("\\", "/")
        parts = name.split("/")
        if name.startswith("/") or (len(name) > 1 and name[1] == ":") or ".." in parts:
            return "unsafe path"
        if info.flag_bits & 0x1:
            return "encrypted"
        if (info.external_attr >> 16) & 0o170000 == 0o120000:
            return "symbolic link"
        ratio = info.file_size / max(info.compress_size, 1)
        if info.file_size > 1024 * 1024 and ratio > settings.ZIP_MAX_COMPRESSION_RATIO:
            raise UnsafeArchiveError(
                f"Archive rejected: '{os.path.basename(name)}' is compressed "
                f"{int(ratio)}:1, over the {settings.ZIP_MAX_COMPRESSION_RATIO}:1 "
                "limit -- this is how zip bombs are built."
            )
        return None

    @staticmethod
    def _copy_member(
        archive: zipfile.ZipFile, info: zipfile.ZipInfo, destination: str, budget: "_ZipBudget"
    ) -> None:
        """Copy one member out, counting every byte against the budget."""
        written = 0
        with archive.open(info) as source, open(destination, "wb") as target:
            while chunk := source.read(1024 * 1024):
                written += len(chunk)
                if written > info.file_size:
                    raise UnsafeArchiveError(
                        f"Archive rejected: '{os.path.basename(info.filename)}' "
                        "is larger than its header declares."
                    )
                budget.take_bytes(len(chunk))
                target.write(chunk)

    @staticmethod
    def _remove_tree(directory: str) -> None:
        """Delete a directory and everything inside it."""
        try:
            for root, dirs, files in os.walk(directory, topdown=False):
                for name in files:
                    try:
                        os.remove(os.path.join(root, name))
                    except OSError:
                        pass
                for name in dirs:
                    try:
                        os.rmdir(os.path.join(root, name))
                    except OSError:
                        pass
            os.rmdir(directory)
        except OSError:
            pass
